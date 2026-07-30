#!/usr/bin/env python3
"""Génère workflows.pptx : 4 schémas de workflow (analyse/implement × GitHub/Jira).
Charte DSFR. Lancer avec le venv : /tmp/pptx-venv/bin/python build_diagrams.py"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

# --- Charte DSFR ---
BLEU  = RGBColor(0x00, 0x00, 0x91)
PERI  = RGBColor(0x6A, 0x6A, 0xF4)
CARD  = RGBColor(0xEC, 0xEC, 0xFE)
LAV   = RGBColor(0xE3, 0xE3, 0xFD)
ROUGE = RGBColor(0xE1, 0x00, 0x0F)
GH    = RGBColor(0x2B, 0x31, 0x37)
JIRA  = RGBColor(0x00, 0x52, 0xCC)
TXT   = RGBColor(0x16, 0x16, 0x16)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE  = RGBColor(0x9A, 0x9A, 0xB8)
FONT  = "Marianne"

KIND = {  # fill, font color, has light border
    "skill":  (BLEU,  WHITE, False),
    "agent":  (PERI,  WHITE, False),
    "sub":    (CARD,  BLEU,  True),
    "gate":   (ROUGE, WHITE, False),
    "github": (GH,    WHITE, False),
    "jira":   (JIRA,  WHITE, False),
    "label":  (LAV,   BLEU,  True),
}

def box(slide, x, y, w, h, kind, title, sub=None, tsize=10, ssize=8):
    fill, fc, border = KIND[kind]
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sp.adjustments[0] = 0.10
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if border:
        sp.line.color.rgb = BLEU; sp.line.width = Pt(0.75)
    else:
        sp.line.color.rgb = fill; sp.line.width = Pt(0.25)
    sp.shadow.inherit = False
    tf = sp.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right"):
        setattr(tf, m, Inches(0.06))
    for m in ("margin_top", "margin_bottom"):
        setattr(tf, m, Inches(0.02))
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = title
    r.font.size = Pt(tsize); r.font.bold = True; r.font.name = FONT
    r.font.color.rgb = fc
    if sub:
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run(); r2.text = sub
        r2.font.size = Pt(ssize); r2.font.name = FONT
        r2.font.color.rgb = fc
    return sp

def arrow(slide, x1, y1, x2, y2, color=BLEU, w=1.5, dash=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                   Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(w)
    c.shadow.inherit = False
    ln = c.line._get_or_add_ln()
    if dash:
        pd = ln.makeelement(qn("a:prstDash"), {"val": dash}); ln.append(pd)
    te = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(te)
    return c

def text(slide, x, y, w, h, s, size=11, bold=True, color=BLEU, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = s
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT
    r.font.color.rgb = color
    return tb

def stack(slide, items, x, y, w, rh, gap):
    """Pile verticale de boîtes + flèches entre elles. items: (kind,title,sub)."""
    cx = x + w / 2
    centers = []
    cy = y
    for i, (k, t, s) in enumerate(items):
        box(slide, x, cy, w, rh, k, t, s)
        centers.append((cx, cy, cy + rh))
        if i > 0:
            arrow(slide, cx, centers[i-1][2] + 0.01, cx, cy - 0.01)
        cy += rh + gap
    return centers

prs = Presentation()
prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

# ============ SLIDE 1 — Titre + légende ============
s = prs.slides.add_slide(BLANK)
text(s, 0.7, 0.55, 12, 1.0, "Pipeline IA de dev — Workflows", size=40, color=BLEU)
text(s, 0.72, 1.5, 12, 0.5,
     "/analyse & /implement  —  egapro (GitHub) et portage Jira (MCP)",
     size=18, bold=False, color=TXT)
text(s, 0.72, 2.45, 6, 0.4, "Légende", size=16, color=BLEU)
leg = [("skill", "Skill / orchestrateur"), ("agent", "Agent (Opus / Sonnet)"),
       ("sub", "Sous-agent · artefact produit"), ("gate", "Gate utilisateur / décision"),
       ("label", "Mode / branche"), ("github", "Point de contact GitHub"),
       ("jira", "Point de contact Jira (MCP / REST)")]
ly = 3.0
for i, (k, lbl) in enumerate(leg):
    col = i // 4; row = i % 4
    bx = 0.75 + col * 6.2; by = ly + row * 0.62
    box(s, bx, by, 0.55, 0.4, k, "")
    text(s, bx + 0.7, by - 0.02, 5.3, 0.45, lbl, size=12, bold=False, color=TXT)
text(s, 0.75, 5.9, 12, 0.4,
     "Conventions :  → flux d'exécution   ·   flèche rouge = régression / boucle de reprise   ·   - - - lecture (read-only)",
     size=12, bold=False, color=TXT)
box(s, 0.75, 6.45, 11.8, 0.7, "gate",
    "Invariant : /analyse ne bouge JAMAIS le board (read-only). "
    "code-dev passe les tickets « In progress ». « In review » / « Done » = humain uniquement.",
    tsize=11.5)

# ============ Générateurs paramétrés ============
def slide_analyse(tracker):
    s = prs.slides.add_slide(BLANK)
    gh = tracker == "github"
    num = "①" if gh else "③"
    name = "egapro (GitHub)" if gh else "projet Jira (MCP)"
    text(s, 0.4, 0.28, 9, 0.7, f"{num}  /analyse — {name}", size=25, color=BLEU)
    box(s, 9.95, 0.34, 2.98, 0.56, "gate",
        "READ-ONLY board" if gh else "READ-ONLY workflow",
        "aucune transition de statut", tsize=11, ssize=8)
    # entrée + détection
    box(s, 4.9, 1.12, 3.55, 0.5, "skill", "/analyse [issue# | description]", tsize=11)
    box(s, 4.55, 1.92, 4.25, 0.62, "skill", "Step 0 — Détection du mode",
        "type d'issue ou mots-clés du prompt", tsize=11)
    arrow(s, 6.67, 1.63, 6.67, 1.91)
    if gh:
        box(s, 9.35, 1.92, 3.55, 0.62, "github", "gh issue view → issueType",
            "Feature · Task · Bug", tsize=10)
    else:
        box(s, 9.35, 1.92, 3.55, 0.62, "jira", "getJiraIssue → issuetype",
            "Epic→epic · Story/Task→task · Bug→bug", tsize=10, ssize=7.5)
    arrow(s, 8.81, 2.23, 9.34, 2.23, color=LINE, w=1.25, dash="dash")
    # 3 colonnes
    colw = 4.0; xs = [0.4, 4.67, 8.93]; top = 2.95; rh = 0.46; gap = 0.10
    ad = ("addCommentToJiraIssue" if not gh else "gh issue comment")
    epic = [
        ("label", "EPIC — type " + ("Feature" if gh else "Epic"), None),
        ("agent", "product-owner (Opus)", "epic body · ## Besoin métier · ## Analyse PO"),
        ("gate", "Gate utilisateur", "Epic validé"),
        ("agent", "architect epic-* (Opus)", "découpe → sous-issues + scénarios"),
        ("gate", "Gate utilisateur", "Architecture validée"),
        ("sub", "Epic + N sous-issues", "body = spec (ticket-spec-format)"),
    ]
    task = [
        ("label", "TASK — type " + ("Task" if gh else "Story/Task"), None),
        ("agent", "architect — mode task (Opus)", "lit body + commentaires · pose des questions si flou"),
        ("gate", "Gate utilisateur", "Analyse validée"),
        ("sub", "commentaire ## Analyse architecte", f"body intact · {ad}"),
    ]
    bug = [
        ("label", "BUG — type Bug", None),
        ("agent", "bug-analyst (Opus)", "repro local / env k8s / visual Figma → root cause"),
        ("gate", "Gate utilisateur", "Analyse validée"),
        ("sub", "commentaire ## Analyse du bug", f"body intact · {ad}"),
    ]
    cE = stack(s, epic, xs[0], top, colw, rh, gap)
    cT = stack(s, task, xs[1], top, colw, rh, gap)
    cB = stack(s, bug, xs[2], top, colw, rh, gap)
    # flèches détection -> têtes de colonnes
    for c in (cE, cT, cB):
        arrow(s, 6.67, 2.55, c[0][0], top - 0.01)
    # barre basse : sizing + report
    by = 6.40
    size_txt = ("Sizing XS→XL des feuilles (set_ticket_size.sh)"
                if gh else "Sizing → editJiraIssue (story points)")
    box(s, 0.4, by, 12.53, 0.48, "skill",
        f"{size_txt}    →    Report · Next : /implement <N>", tsize=11.5)
    for c in (cE, cT, cB):
        arrow(s, c[0][0], c[-1][2] + 0.01, c[0][0], by - 0.01)
    if not gh:
        text(s, 0.4, 6.98, 12.5, 0.3,
             "Outils MCP : getJiraIssue · createJiraIssue · addCommentToJiraIssue · editJiraIssue   "
             "(noms = serveur Atlassian Rovo ; vérifier sur le serveur installé)",
             size=8, bold=False, color=TXT)

def slide_implement(tracker):
    s = prs.slides.add_slide(BLANK)
    gh = tracker == "github"
    num = "②" if gh else "④"
    name = "egapro (GitHub)" if gh else "projet Jira (MCP + REST)"
    text(s, 0.4, 0.28, 9, 0.7, f"{num}  /implement — {name}", size=25, color=BLEU)
    # entrée + steps
    box(s, 4.85, 0.98, 3.6, 0.46, "skill", "/implement <issue#>", tsize=11)
    box(s, 3.95, 1.58, 5.4, 0.56, "skill", "Step 0–1 — type d'issue + analyse présente ?",
        "## Analyse PO/architecte/bug · subIssues>0 — sinon → /analyse", tsize=10.5, ssize=8)
    arrow(s, 6.65, 1.45, 6.65, 1.57)
    if not gh:
        box(s, 9.55, 0.98, 3.38, 1.18, "gate", "⚠  MCP ≠ bash",
            "Phase conversationnelle (task/bug) → Jira MCP.  "
            "Boucle epic (nohup, sans LLM) → Jira REST via curl.", tsize=11, ssize=8)
    # 2 colonnes
    top = 2.78; rh = 0.46; gap = 0.10
    xL = 0.4; xR = 6.93; w = 6.0
    epic = [
        ("label", "EPIC (Feature)", "board via Jira REST (curl)" if not gh else "branche d'intégration epic/<N>"),
        ("skill", "Aligner main worktree → epic/<N>", "ensure_epic_branch.sh · refuse si dirty"),
        ("skill", "nohup epic_loop.sh & — background",
         "dispatch parallèle, idempotent" if gh else "mutations board en REST /transitions"),
        ("agent", "code-dev (main agent) ×N parallèle", "1 worktree/ticket · spawne 6 sous-agents ↓"),
        ("skill", "squash-merge PR → epic/<N>", "process_tick_result.sh · débloque les dépendances"),
        ("agent", "Gate E2E (e2e-dev) → doc-writer → PR finale", "régression → architect-rework → re-loop"),
        ("gate", "Gate d'acceptation utilisateur (/loop /report)", "changements → architect-rework (boucle)"),
    ]
    taskbug = [
        ("label", "TASK / BUG", "board via Jira MCP" if not gh else "PR mergée à la main (humain)"),
        ("skill", "worktree + docker · create_linked_branch",
         "statut → In progress" + ("" if gh else " (transitionJiraIssue)")),
        ("agent", "code-dev (CLI foreground, main agent)", "Sonnet|Opus · spawne 6 sous-agents ↓"),
        ("gate", "Verdict JSON", "validated→e2e-dev · refacto→/analyse · regression→architect-rework"),
        ("agent", "e2e-dev (Opus)", "couverture E2E · push sur la PR (re-déclenche la CI)"),
        ("sub", "PR draft → ready", "ticket reste In progress"),
    ]
    cL = stack(s, epic, xL, top, w, rh, gap)
    cR = stack(s, taskbug, xR, top, w, rh, gap)
    arrow(s, 6.65, 2.14, cL[0][0], top - 0.01)
    arrow(s, 6.65, 2.14, cR[0][0], top - 0.01)
    # barre basse : cluster sous-agents
    by = 6.72
    box(s, 0.4, by, 12.53, 0.52, "sub",
        "▸ Sous-agents spawnés par code-dev (dans les deux modes)",
        "tu-dev (TU + intégration) · validator · structural-auditor · rgaa-auditor · security-auditor · functional-validator",
        tsize=10.5, ssize=9)
    foot = ("code-dev = main agent (un sous-agent ne peut pas en spawner). "
            "Board : In progress = code-dev · In review / Done = humain.")
    if not gh:
        foot = ("Code / PR restent sur l'hôte git ; lien Jira via clé d'issue (Smart Commits). "
                + foot)
    text(s, 0.4, 7.32, 12.5, 0.3, foot, size=8, bold=False, color=TXT)

slide_analyse("github")
slide_implement("github")
slide_analyse("jira")
slide_implement("jira")

out = "/home/selim/Documents/clients/ministeres-sociaux/egapro/docs/tuto-config/workflows.pptx"
prs.save(out)
print("saved", out, "—", len(prs.slides._sldIdLst), "slides")
