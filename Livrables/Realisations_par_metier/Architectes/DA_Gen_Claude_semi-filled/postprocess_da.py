#!/usr/bin/env python3
"""postprocess_da.py — complete un DA genere par fill_da.py.

Le gabarit MASS a une capacite FIXE : 2 blocs serveurs (cadre 9, 4 emplacements)
et 5 blocs URL (cadre 12). Quand contenu.json en contient davantage, fill_da.py
ecrit ce qui rentre et laisse tomber le reste (silencieusement).

Ce script rattrape ce debordement : il clone des blocs VIDES depuis le gabarit
et y ecrit les entrees `serveurs`/`urls` de contenu.json absentes du document.
100% local, aucune dependance reseau.

Usage:
    python postprocess_da.py DA_Modele.docx contenu.json entree.docx sortie.docx

A relancer si le .docx est regenere par fill_da.py (ces blocs sont post-traitement).
"""
import sys, copy, json, os
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import fill_da as F  # reutilise les helpers du pipeline (set_cell, _unique_cells...)


def collect(document):
    res = []
    def walk(tables):
        for t in tables:
            res.append(t)
            for r in t.rows:
                for c in r.cells:
                    walk(c.tables)
    walk(document.tables)
    return res


def uniq(document, anchor):
    """Blocs physiques distincts contenant l'ancre (dedup par element _tbl)."""
    seen, out = set(), []
    for t in collect(document):
        if F.find_in_table(t, anchor) and t._tbl not in seen:
            seen.add(t._tbl)
            out.append(t)
    return out


def clone_after(anchor_tbl, src_tbl):
    """Insere une copie de src_tbl apres anchor_tbl, avec paragraphes separateurs
    (deux <w:tbl> accoles seraient fusionnes par Word)."""
    new = copy.deepcopy(src_tbl)
    nxt = anchor_tbl.getnext()
    p_new = OxmlElement("w:p")
    if nxt is not None and nxt.tag == qn("w:p"):
        nxt.addnext(new)        # anchor, p_sep, new, ...
        new.addnext(p_new)      # anchor, p_sep, new, p_new, ...
    else:
        anchor_tbl.addnext(new)
        new.addprevious(OxmlElement("w:p"))
        new.addnext(p_new)
    return new


def server_names(block):
    cells = F._unique_cells(block.rows[0])
    pos = [i for i, c in enumerate(cells) if "Nom serveur" in F.ctext(c)]
    return [(F.ctext(cells[p + 1]) if p + 1 < len(cells) else "").strip() for p in pos]


def url_label(block):
    return (F.ctext(F._unique_cells(block.rows[0])[-1]) or "").strip()


def main():
    model, contenu, src, out = sys.argv[1:5]
    data = json.load(open(contenu, encoding="utf-8"))

    tpl = Document(model)
    tpl_srv = uniq(tpl, "Nom serveur (logique)")[0]._tbl   # bloc serveur vide
    tpl_url = uniq(tpl, "Libellé URL")[0]._tbl              # bloc URL vide

    doc = Document(src)

    # --- serveurs manquants ---
    srv_blocks = uniq(doc, "Nom serveur (logique)")
    present = {n for b in srv_blocks for n in server_names(b) if n}
    srv_missing = [s for s in data.get("serveurs", []) if s["nom"] not in present]
    n_srv_blocks = (len(srv_missing) + 1) // 2  # 2 emplacements par bloc

    # --- urls manquantes ---
    url_blocks = uniq(doc, "Libellé URL")
    present_u = {url_label(b) for b in url_blocks if url_label(b)}
    url_missing = [u for u in data.get("urls", []) if u["libelle"] not in present_u]

    # Phase A : inserer les blocs vides
    if n_srv_blocks:
        anchor = srv_blocks[-1]._tbl
        for _ in range(n_srv_blocks):
            anchor = clone_after(anchor, tpl_srv)
    if url_missing:
        anchor = uniq(doc, "Libellé URL")[-1]._tbl
        for _ in range(len(url_missing)):
            anchor = clone_after(anchor, tpl_url)

    tmp = out + ".tmp.docx"
    doc.save(tmp)

    # Phase B : reouverture propre + remplissage des blocs vides
    doc = Document(tmp)
    it = iter(srv_missing)
    for b in uniq(doc, "Nom serveur (logique)"):
        cells = F._unique_cells(b.rows[0])
        pos = [i for i, c in enumerate(cells) if "Nom serveur" in F.ctext(c)]
        if all(not n for n in server_names(b)):  # bloc vide ajoute
            for p in pos:
                s = next(it, None)
                if s is None:
                    break
                F.set_cell(cells[p + 1], s["nom"])
                F.fill_components(b, p, s.get("composants", []))
    itu = iter(url_missing)
    for b in uniq(doc, "Libellé URL"):
        if not url_label(b):  # bloc vide ajoute
            u = next(itu, None)
            if u is None:
                break
            F.set_cell(F._unique_cells(b.rows[0])[-1], u["libelle"])
            if len(b.rows) >= 3:
                cells = F._unique_cells(b.rows[2])
                for ci, key in enumerate(["acteur", "ressource", "fonctionnalite", "donnees"]):
                    if ci < len(cells):
                        F.set_cell(cells[ci], u.get(key))
    doc.save(out)
    os.remove(tmp)
    print(f"OK — {len(srv_missing)} serveur(s) et {len(url_missing)} URL(s) ajoutes -> {out}")


if __name__ == "__main__":
    main()
