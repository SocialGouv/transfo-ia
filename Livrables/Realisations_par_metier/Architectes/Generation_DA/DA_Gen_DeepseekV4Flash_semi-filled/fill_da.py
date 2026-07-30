#!/usr/bin/env python3
"""
fill_da.py — remplit une copie de DA_Modele.docx a partir d'un objet JSON de
contenu produit par le prompt de generation (voir prompt_generation_DA.md).

Usage:
    python fill_da.py DA_Modele.docx contenu.json DA-Projet-V0.1.0.docx
    python fill_da.py DA_Modele.docx contenu.json sortie.docx --schemas=render
    python fill_da.py DA_Modele.docx contenu.json sortie.docx --schemas=placeholder
    python fill_da.py DA_Modele.docx contenu.json sortie.docx --schemas=none

Principe: reperage par ANCRE (texte de libelle), jamais par index absolu.
Les grilles dynamiques (fonctionnalites, donnees, composants, flux, URLs)
grossissent par clonage de ligne.

SCHEMAS (cadres 5-8) — nouveau. Le gabarit MASS n'est PAS vierge : il embarque
les 4 schemas d'EXEMPLE (projet « Transparence Sante », filigranes « Exemple »)
comme objets OLE. Sans traitement, ils sont recopies tels quels dans chaque DA
genere (bug historique : « les schemas restent des schemas d'exemple »). Ce
script neutralise donc systematiquement ces 4 objets d'exemple et, selon le mode
--schemas :
 - render (defaut) : rend les diagrammes Mermaid du bloc JSON `schemas` en PNG
   (mermaid-cli, rendu 100% LOCAL — rien n'est envoye a l'exterieur) et les
   insere a la place des exemples, avec un filigrane « brouillon a redessiner
   en palette MASS ». Degradation gracieuse -> placeholder si le rendu echoue.
 - placeholder : remplace l'exemple par un encart « Schema a produire », sans
   image. La specification structuree + la source Mermaid partent en annexe.
 - none : retire simplement l'exemple (zone vide).
Dans tous les cas, la specification de chaque schema et sa source Mermaid sont
ecrites dans une annexe « Specifications de schemas » creee en fin de document.
La legende generique « Niveaux de securite MASS » (image2) est preservee.

Le script est volontairement defensif: tout cadre/objet introuvable est signale
en avertissement et saute, il ne casse jamais le document.

Limites connues (a completer/verifier manuellement dans le .docx genere):
 - Cadres 5-8 : le rendu Mermaid est un BROUILLON fidele au code, pas le dessin
   final en palette MASS (travail manuel, cf. la pptx « copier ajuster coller »).
 - Cadre 9: le nom du serveur et les composants sont remplis; Type / Role /
   VCPU / RAM restent a saisir (cellules fusionnees non ecrites volontairement).
 - Cadre 3 « Sensibilite des donnees » (grille de cases) et blocs utilisabilite
   tablette/smartphone: cases a cocher, non automatises.
 - Cadre 4 (DICT, EBIOS, periodes, garantie de service, temps de reponse) et
   cadre 11 (dimensionnement): champs « humains », a remplir via le questionnaire.
 - Annexe de suivi des versions: a renseigner (1re ligne V0.1.0).
Le rapport de provenance et le questionnaire de completion sont produits par le
modele (voir prompt_generation_DA.md), pas par ce script.
"""
import copy
import glob
import json
import os
import re
import subprocess
import sys
import tempfile
from docx import Document
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

WARN = []
INFO = []

# namespaces VML / OLE (les schemas d'exemple sont des objets OLE PBrush)
NS_V = "urn:schemas-microsoft-com:vml"
NS_O = "urn:schemas-microsoft-com:office:office"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def warn(msg):
    WARN.append(msg)


def info(msg):
    INFO.append(msg)


# ---------- navigation ----------
def all_tables(doc):
    """Tous les tableaux, y compris imbriques, en profondeur."""
    out = []

    def rec(tbls):
        for t in tbls:
            out.append(t)
            for row in t.rows:
                for c in _unique_cells(row):
                    rec(c.tables)
    rec(doc.tables)
    return out


def _unique_cells(row):
    seen, cells = set(), []
    for c in row.cells:
        k = id(c._tc)
        if k in seen:
            continue
        seen.add(k)
        cells.append(c)
    return cells


def ctext(cell):
    return " ".join(p.text for p in cell.paragraphs).strip()


def find_table(tables, anchor, contains=True):
    """Premier tableau dont une cellule (n'importe ou) matche l'ancre."""
    for t in tables:
        for row in t.rows:
            for c in _unique_cells(row):
                txt = ctext(c)
                if (anchor in txt) if contains else (anchor == txt):
                    return t
    warn(f"ancre introuvable: {anchor!r}")
    return None


def set_cell(cell, text):
    if text is None:
        return
    # ecrit dans le 1er paragraphe, vide les suivants, conserve le style
    paras = cell.paragraphs
    p0 = paras[0]
    for r in list(p0.runs):
        r.text = ""
    if p0.runs:
        p0.runs[0].text = str(text)
    else:
        p0.add_run(str(text))
    for extra in paras[1:]:
        for r in list(extra.runs):
            r.text = ""


def clone_row_after(table, row):
    """Clone une ligne (structure/style) et l'insere juste apres elle."""
    new_tr = copy.deepcopy(row._tr)
    row._tr.addnext(new_tr)
    # nettoie le texte des cellules clonees
    for tc in new_tr.findall(qn("w:tc")):
        for t in tc.findall(".//" + qn("w:t")):
            t.text = ""
    return table.rows[[r._tr for r in table.rows].index(new_tr)]


# ---------- helpers metier ----------
def mrep_columns(header_row):
    """Indices (dans les cellules uniques) des colonnes M/R/E/P."""
    cols = {}
    for i, c in enumerate(_unique_cells(header_row)):
        t = ctext(c)
        if t in ("M", "R", "E", "P"):
            cols[t] = i
    return cols


def fill_labeled_block(tables, anchor, value):
    """Bloc mono-colonne: ancre en tete, valeur dans la 1re ligne suivante."""
    t = find_table(tables, anchor)
    if t is None or value is None:
        return
    if len(t.rows) >= 2:
        set_cell(_unique_cells(t.rows[1])[0], value)


def fill_keyvalue(tables, anchor, code, value):
    """Ligne dont la 1re cellule == code (ex D1); valeur en derniere cellule."""
    t = find_table(tables, anchor)
    if t is None:
        return
    for row in t.rows:
        cells = _unique_cells(row)
        if cells and ctext(cells[0]) == code:
            set_cell(cells[-1], value)
            return
    warn(f"code {code!r} introuvable sous {anchor!r}")


def fill_mrep_grid(tables, anchor, items, label_col=0, total_kw="Total",
                   extra_cols=None, fixed=False):
    """Remplit une grille (libelle + colonnes M/R/E/P), lignes clonees au besoin.
    items: [{'libelle':..., 'mrep':[...], <extra>:...}]
    extra_cols: dict {json_key: cell_index} pour colonnes intermediaires.
    fixed=True: table a lignes fixes (ex. Services utilises) -> on repere la
    ligne existante par son libelle et on la coche, sans inserer de ligne.
    """
    t = find_table(tables, anchor)
    if t is None:
        return
    # localise la ligne d'en-tete portant M/R/E/P
    header_row, cols = None, {}
    for row in t.rows:
        c = mrep_columns(row)
        if len(c) >= 3:
            header_row, cols = row, c
            break
    if header_row is None:
        warn(f"pas d'en-tete M/R/E/P sous {anchor!r}")
        return
    header_idx = [r._tr for r in t.rows].index(header_row._tr)
    if fixed:
        for item in items:
            lib = (item.get("libelle") or "").lower()
            matched = None
            for row in t.rows[header_idx + 1:]:
                cells = _unique_cells(row)
                if cells and lib and lib in ctext(cells[label_col]).lower():
                    matched = cells
                    break
            if matched is None:
                warn(f"ligne fixe introuvable pour {item.get('libelle')!r} "
                     f"sous {anchor!r}")
                continue
            for key, ci in (extra_cols or {}).items():
                if item.get(key) and ci < len(matched):
                    set_cell(matched[ci], item[key])
            for m in item.get("mrep", []):
                if m in cols and cols[m] < len(matched):
                    set_cell(matched[cols[m]], "x")
        return
    # ligne total (a garder en bas) et 1re ligne de saisie
    total_idx = None
    for i, row in enumerate(t.rows):
        if total_kw.lower() in ctext(_unique_cells(row)[0]).lower():
            total_idx = i
    first_data = header_idx + 1
    n_slots = (total_idx - first_data) if total_idx else (len(t.rows) - first_data)
    # cree assez de lignes en clonant la 1re ligne de saisie
    while n_slots < len(items):
        clone_row_after(t, t.rows[first_data])
        n_slots += 1
    for j, item in enumerate(items):
        row = t.rows[first_data + j]
        cells = _unique_cells(row)
        set_cell(cells[label_col], item.get("libelle", ""))
        for key, ci in (extra_cols or {}).items():
            if item.get(key) is not None and ci < len(cells):
                set_cell(cells[ci], item[key])
        for m in item.get("mrep", []):
            if m in cols and cols[m] < len(cells):
                set_cell(cells[cols[m]], "x")
    # recalcule les totaux si presents
    if total_idx is not None:
        trow = _unique_cells(t.rows[[r._tr for r in t.rows].index(t.rows[total_idx]._tr)])
        for m, ci in cols.items():
            n = sum(1 for it in items if m in it.get("mrep", []))
            if ci < len(trow):
                set_cell(trow[ci], str(n))


def fill_rowset(tables, anchor, header_kw, items, keys):
    """Table type matrice de flux / URLs blocs / planning: en-tete + lignes.
    keys: liste des cles json alignees sur les colonnes (cellules uniques).
    """
    t = find_table(tables, anchor)
    if t is None:
        return
    # ligne d'en-tete
    hidx = None
    for i, row in enumerate(t.rows):
        if header_kw in ctext(_unique_cells(row)[0]):
            hidx = i
            break
    if hidx is None:
        hidx = 0
    first = hidx + 1
    slots = len(t.rows) - first
    while slots < len(items):
        clone_row_after(t, t.rows[first])
        slots += 1
    for j, item in enumerate(items):
        cells = _unique_cells(t.rows[first + j])
        for ci, key in enumerate(keys):
            if key and item.get(key) is not None and ci < len(cells):
                set_cell(cells[ci], item[key])


def compute_volumes(v, kind):
    """Calcule D5/D6/D7 ou F5/F6/F7 si les entrees sont numeriques."""
    def num(x):
        try:
            return float(str(x).replace(",", "."))
        except Exception:
            return None
    p = "D" if kind == "d" else "F"
    a, b, c, d = (num(v.get(p + "1")), num(v.get(p + "2")),
                  num(v.get(p + "3")), num(v.get(p + "4")))
    div = 1024 * 1024 if kind == "d" else 1024
    out = {}
    if a is not None and c is not None:
        out[p + "5"] = round(a * c / div, 2)
    if b is not None and c is not None:
        out[p + "6"] = round(b * c / div, 2)
    if d is not None and out.get(p + "6") is not None and out.get(p + "5") is not None:
        out[p + "7"] = round(d * out[p + "6"] + out[p + "5"], 2)
    return out


# ---------- schemas (cadres 5-8) ----------
CADRES = ["cadre5", "cadre6", "cadre7", "cadre8"]
CADRE_TITRES = {
    "cadre5": "Cadre 5 — Architecture Acteurs",
    "cadre6": "Cadre 6 — Architecture Fonctionnelle",
    "cadre7": "Cadre 7 — Architecture Applicative",
    "cadre8": "Cadre 8 — Architecture technique",
}


def _style_dims_pt(shape_el):
    """Largeur/hauteur (en points) lues dans l'attribut style d'une <v:shape>."""
    style = (shape_el.get("style") or "") if shape_el is not None else ""
    w = h = None
    m = re.search(r"width:([0-9.]+)pt", style)
    if m:
        w = float(m.group(1))
    m = re.search(r"height:([0-9.]+)pt", style)
    if m:
        h = float(m.group(1))
    return w, h


def find_schema_objects(doc):
    """Retourne les <w:object> des 4 schemas (cadres 5-8), dans l'ordre du
    document, en excluant la legende generique.

    Heuristique robuste (independante des noms de media) : on ecarte les objets
    au ratio tres large (legende « Niveaux de securite », ratio > 2.2) et on
    prend les objets restants dans l'ordre. On valide qu'il en reste bien 4.
    """
    body = doc.element.body
    schema = []
    for ob in body.iter(qn("w:object")):
        shape = ob.find(".//{%s}shape" % NS_V)
        w, h = _style_dims_pt(shape)
        ratio = (w / h) if (w and h) else 0
        is_legend = ratio > 2.2  # la legende MASS est nettement plus large que haute
        if is_legend:
            info(f"objet ignore (legende presumee, ratio={ratio:.2f})")
            continue
        schema.append((ob, w, h))
    if len(schema) != 4:
        warn(f"{len(schema)} objet(s) schema detecte(s) au lieu de 4 : "
             f"le mapping cadre 5-8 peut etre imprecis, verifier le .docx.")
    return schema[:4]


def render_mermaid(src, out_png, chrome=None, scale=3, width=1400, timeout=180):
    """Rend une source Mermaid en PNG, 100% LOCAL via mermaid-cli (mmdc).
    Retourne True si le PNG a bien ete produit, False sinon (jamais d'exception)."""
    if not src or not src.strip():
        return False
    tmpd = tempfile.mkdtemp(prefix="da_mmd_")
    mmd = os.path.join(tmpd, "in.mmd")
    cfg = os.path.join(tmpd, "puppeteer.json")
    with open(mmd, "w", encoding="utf-8") as f:
        f.write(src)
    with open(cfg, "w", encoding="utf-8") as f:
        json.dump({"args": ["--no-sandbox", "--disable-dev-shm-usage",
                            "--disable-gpu"]}, f)
    env = dict(os.environ)
    if chrome:
        env["PUPPETEER_EXECUTABLE_PATH"] = chrome
    cmd = ["npx", "-y", "@mermaid-js/mermaid-cli@11",
           "-i", mmd, "-o", out_png, "-b", "white",
           "-w", str(width), "-s", str(scale),
           "--puppeteerConfigFile", cfg]
    try:
        r = subprocess.run(cmd, env=env, capture_output=True, text=True,
                           timeout=timeout)
        if r.returncode != 0 or not os.path.exists(out_png):
            warn("rendu Mermaid echoue: " +
                 (r.stderr.strip().splitlines()[-1] if r.stderr.strip() else
                  f"code {r.returncode}"))
            return False
        return True
    except FileNotFoundError:
        warn("rendu Mermaid indisponible: 'npx' introuvable (Node.js requis).")
        return False
    except subprocess.TimeoutExpired:
        warn(f"rendu Mermaid : delai depasse ({timeout}s).")
        return False
    except Exception as e:  # ceinture + bretelles : ne jamais casser la generation
        warn(f"rendu Mermaid : erreur {e}")
        return False


def find_chrome(auto_install=True):
    """Localise un Chrome/Chromium NON-snap (le snap est confine par AppArmor et
    ne peut pas lire l'index.html de mermaid-cli -> ERR_ACCESS_DENIED)."""
    for e in ("DA_CHROME", "PUPPETEER_EXECUTABLE_PATH"):
        p = os.environ.get(e)
        if p and os.path.exists(p) and "snap" not in os.path.realpath(p):
            return p
    cands = sorted(glob.glob(os.path.expanduser(
        "~/.cache/puppeteer/chrome/*/chrome-linux*/chrome")))
    if cands:
        return cands[-1]
    for p in ("/usr/bin/google-chrome", "/usr/bin/google-chrome-stable",
              "/opt/google/chrome/chrome"):
        if os.path.exists(p) and "snap" not in os.path.realpath(p):
            return p
    if auto_install:
        info("Aucun Chrome non-snap trouve : installation via puppeteer "
             "(1re fois, ~150 Mo)...")
        try:
            subprocess.run(["npx", "-y", "puppeteer@24", "browsers", "install",
                            "chrome"], capture_output=True, text=True, timeout=600)
        except Exception as e:
            warn(f"installation Chrome echouee: {e}")
        cands = sorted(glob.glob(os.path.expanduser(
            "~/.cache/puppeteer/chrome/*/chrome-linux*/chrome")))
        if cands:
            return cands[-1]
    warn("Aucun Chrome exploitable pour le rendu Mermaid (le chromium snap ne "
         "convient pas). Bascule en mode placeholder.")
    return None


def _clear_run(run_el):
    """Vide un <w:r> de ses enfants (retire l'objet OLE d'exemple)."""
    for child in list(run_el):
        run_el.remove(child)


def _para_of(ob):
    run = ob.getparent()
    para = run.getparent() if run is not None else None
    return run, para


def _picture_kwargs(png_path, box_w_pt, box_h_pt):
    """Dimensionne l'image pour tenir dans la largeur du cadre, avec un plafond
    de hauteur (évite qu'un diagramme très haut déborde sur plusieurs pages).
    PIL est optionnel : sans lui, on contraint simplement la largeur."""
    box_w = box_w_pt or 471.0
    max_h = 620.0  # ~ hauteur utile d'une page A4
    try:
        from PIL import Image
        iw, ih = Image.open(png_path).size
        if iw and ih:
            h = box_w * ih / iw
            if h > max_h:
                return {"height": Pt(max_h)}
    except Exception:
        pass
    return {"width": Pt(box_w)}


def replace_schema_object(doc, ob, w_pt, h_pt, mode, png_path, caption):
    """Remplace un objet OLE d'exemple par une image rendue (mode=render) ou un
    encart texte (placeholder/none)."""
    run, para_el = _para_of(ob)
    if run is None or para_el is None or para_el.tag != qn("w:p"):
        warn("objet schema sans paragraphe parent : ignore.")
        return
    _clear_run(run)  # retire l'objet OLE (image + embedding d'exemple)
    p = Paragraph(para_el, doc)
    if mode == "render" and png_path and os.path.exists(png_path):
        r_img = p.add_run()
        r_img.add_picture(png_path, **_picture_kwargs(png_path, w_pt, h_pt))
        r_cap = p.add_run()
        r_cap.add_break()
        r_cap2 = p.add_run(caption)
        r_cap2.italic = True
        r_cap2.font.size = Pt(8)
    elif mode == "placeholder":
        r = p.add_run(caption)
        r.italic = True
        r.font.size = Pt(10)
    else:  # none : on laisse la zone vide (l'exemple a ete retire)
        pass


def fill_schemas(doc, schemas, version, mode, chrome):
    """Neutralise les 4 schemas d'exemple et injecte le contenu du projet."""
    objs = find_schema_objects(doc)
    if not objs:
        warn("aucun objet schema localise dans le gabarit : cadres 5-8 non traites.")
        return
    render_dir = tempfile.mkdtemp(prefix="da_render_")
    for i, (ob, w_pt, h_pt) in enumerate(objs):
        key = CADRES[i] if i < len(CADRES) else f"cadre{5 + i}"
        sc = (schemas or {}).get(key, {}) if isinstance(schemas, dict) else {}
        titre = sc.get("titre") or CADRE_TITRES.get(key, key)
        mermaid = sc.get("mermaid", "")
        png = None
        used_mode = mode
        if mode == "render":
            if not mermaid.strip():
                warn(f"{key}: pas de source Mermaid dans le JSON -> placeholder.")
                used_mode = "placeholder"
            else:
                png = os.path.join(render_dir, f"{key}.png")
                if not render_mermaid(mermaid, png, chrome=chrome):
                    used_mode = "placeholder"
                    png = None
        if used_mode == "render":
            cap = (f"Brouillon auto-généré depuis le code source ({version}) "
                   f"— {titre}. À redessiner en palette MASS ; spécification "
                   f"détaillée en annexe « Spécifications de schémas ».")
        else:
            cap = (f"[{titre} — schéma à produire en palette MASS. "
                   f"Spécification et source Mermaid en annexe "
                   f"« Spécifications de schémas ».]")
        replace_schema_object(doc, ob, w_pt, h_pt, used_mode, png, cap)
        info(f"{key}: schema d'exemple remplace (mode={used_mode}).")


def _add_mono(doc, text):
    """Ajoute un paragraphe monospace (source Mermaid) sans dependre d'un style."""
    p = doc.add_paragraph()
    for j, line in enumerate(text.splitlines() or [""]):
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(8)
        if j < len(text.splitlines()) - 1:
            r.add_break()
    return p


def _heading(doc, text, level):
    """Titre d'annexe. On ajoute UN seul paragraphe et on applique le style
    Heading si le gabarit le definit (MASS ne definit pas tous les niveaux),
    sinon on retombe sur du gras dimensionne. Evite le doublon d'add_heading
    (qui insere le texte puis leve sur style absent)."""
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(max(9, 16 - 2 * level))
    style_name = "Title" if level == 0 else f"Heading {level}"
    try:
        p.style = doc.styles[style_name]
    except Exception:
        pass
    return p


def add_schema_annex(doc, schemas, flux):
    """Ajoute l'annexe « Specifications de schemas » (spec + source Mermaid +
    table des flux pour le cadre 8). Cette annexe n'existe pas dans le gabarit
    d'origine : le prompt y renvoyait a tort, on la cree ici."""
    if not isinstance(schemas, dict) or not any(schemas.get(k) for k in CADRES):
        info("annexe schemas non ajoutee (bloc `schemas` absent du JSON).")
        return
    doc.add_page_break()
    _heading(doc, "Annexe — Spécifications de schémas (cadres 5 à 8)", 1)
    doc.add_paragraph(
        "Base de redessin en palette MASS (cf. « DA templates schémas copier "
        "ajuster coller.pptx »). Contenu dérivé du code source : à valider par "
        "l'architecte SI.")
    for key in CADRES:
        sc = schemas.get(key) or {}
        titre = sc.get("titre") or CADRE_TITRES.get(key, key)
        _heading(doc, titre, 2)
        spec = sc.get("specification", "")
        if spec:
            for block in str(spec).split("\n\n"):
                doc.add_paragraph(block.strip())
        if key == "cadre8" and flux:
            _heading(doc, "Matrice des flux (cf. cadre 10)", 3)
            tbl = doc.add_table(rows=1, cols=5)
            try:
                tbl.style = "Table Grid"
            except Exception:
                pass
            hdr = tbl.rows[0].cells
            for c, txt in zip(hdr, ["N°", "Source", "Destination", "Protocole",
                                    "Commentaires"]):
                c.paragraphs[0].add_run(txt).bold = True
            for fx in flux:
                cells = tbl.add_row().cells
                cells[0].text = str(fx.get("n", ""))
                cells[1].text = str(fx.get("source", ""))
                cells[2].text = str(fx.get("destination", ""))
                cells[3].text = str(fx.get("protocole", ""))
                cells[4].text = str(fx.get("commentaires", ""))
        mermaid = sc.get("mermaid", "")
        if mermaid.strip():
            _heading(doc, "Source Mermaid", 3)
            _add_mono(doc, mermaid.strip())


def check_flux_consistency(schemas, flux):
    """Verifie (souple) que les numeros de flux du cadre 8 (schema technique)
    couvrent ceux de la matrice cadre 10, comme l'exige la methode MASS."""
    if not isinstance(schemas, dict):
        return
    mer = (schemas.get("cadre8") or {}).get("mermaid", "")
    if not mer or not flux:
        return
    declared = {str(fx.get("n")).strip() for fx in flux if fx.get("n") is not None}
    # numeros presents dans le schema (jetons numeriques isoles)
    present = set(re.findall(r"(?<!\d)(\d{1,3})(?!\d)", mer))
    missing = sorted(declared - present, key=lambda x: (len(x), x))
    if missing:
        warn("cohérence flux (cadre 8 ↔ cadre 10) : numéros de la matrice "
             f"absents du schema technique : {', '.join(missing)}.")


def _version_from_path(path):
    m = re.search(r"V\d+\.\d+\.\d+(?:\.\d+)?", os.path.basename(path))
    return m.group(0) if m else "V0.1.0"


def prune_orphan_media(doc):
    """Supprime du paquet les media/OLE de la PARTIE DOCUMENT qui ne sont plus
    references (les schemas d'exemple retires). Sûr : on ne coupe une relation
    que si son rId n'apparait plus nulle part dans le corps ; les parts rendues
    inatteignables ne sont pas ecrites a l'enregistrement (comportement OPC).
    La legende et les images rendues (referencees) sont conservees."""
    body = doc.element.body
    used = set()
    for el in body.iter():
        for attr in (qn("r:embed"), qn("r:id"), qn("r:link")):
            v = el.get(attr)
            if v:
                used.add(v)
    dropped = []
    for rid, rel in list(doc.part.rels.items()):
        rt = getattr(rel, "reltype", "")
        if ("image" in rt or "oleObject" in rt) and rid not in used:
            try:
                del doc.part.rels[rid]
                dropped.append(rid)
            except Exception:
                pass
    if dropped:
        info(f"{len(dropped)} media/OLE d'exemple orphelins supprimes du paquet.")


# ---------- pilote ----------
def parse_args(argv):
    pos, mode = [], "render"
    for a in argv:
        if a.startswith("--schemas="):
            mode = a.split("=", 1)[1].strip().lower()
        elif a in ("-h", "--help"):
            print(__doc__)
            sys.exit(0)
        else:
            pos.append(a)
    if mode not in ("render", "placeholder", "none"):
        print(f"Mode --schemas inconnu: {mode!r} (render|placeholder|none)")
        sys.exit(1)
    return pos, mode


def main():
    pos, mode = parse_args(sys.argv[1:])
    if len(pos) != 3:
        print(__doc__)
        sys.exit(1)
    model, data_path, out_path = pos
    data = json.load(open(data_path, encoding="utf-8"))
    doc = Document(model)
    T = all_tables(doc)

    proj = data.get("projet", {})
    fill_labeled_block(T, "Nom du projet applicatif", proj.get("nom"))
    fill_labeled_block(T, "Contexte projet applicatif", proj.get("contexte"))
    fill_labeled_block(T, "Objectifs projet applicatif", proj.get("objectifs"))
    fill_labeled_block(T, "Enjeux projet applicatif", proj.get("enjeux"))

    fill_rowset(T, "Planning projet", "Version",
                data.get("planning", []), ["version", "date", "commentaires"])
    fill_rowset(T, "Acteurs du projet", "Rôle",
                _by_role(T, data.get("acteurs_projet", [])),
                ["role", "nom", "fonction", "entite"])

    fill_mrep_grid(T, "Acteurs métiers du SI applicatif",
                   data.get("acteurs_metiers", []),
                   extra_cols={"nombre": 1})
    fill_mrep_grid(T, "Fonctionnalités du SI applicatif",
                   data.get("fonctionnalites", []))
    fill_mrep_grid(T, "Données métier du SI Applicatif",
                   data.get("donnees_metier", []))
    fill_mrep_grid(T, "Fichiers métiers du SI applicatif",
                   data.get("fichiers_metier", []))
    fill_mrep_grid(T, "Référentiel données", data.get("referentiels", []),
                   extra_cols={"mode": 1})
    fill_mrep_grid(T, "Services utilisés par application",
                   data.get("services_utilises", []), extra_cols={"mode": 1},
                   fixed=True)

    fill_rowset(T, "Dépendances avec d’autres SI", "Système",
                data.get("dependances_si", []),
                ["si", "fournisseur", "consommateur"])

    vd = data.get("volumetrie_donnees", {})
    vd.update(compute_volumes(vd, "d"))
    for k in ("D1", "D2", "D3", "D4", "D5", "D6", "D7"):
        fill_keyvalue(T, "Volumétrie données du SI Applicatif", k, vd.get(k))
    vf = data.get("volumetrie_fichiers", {})
    vf.update(compute_volumes(vf, "f"))
    for k in ("F1", "F2", "F3", "F4", "F5", "F6", "F7"):
        fill_keyvalue(T, "Volumétrie Fichiers du SI Applicatif", k, vf.get(k))

    fill_servers(T, data.get("serveurs", []))
    flux = data.get("flux", [])
    fill_rowset(T, "N° Flux", "N° Flux", flux,
                ["n", "source", "destination", "protocole", "commentaires"])
    fill_urls(T, data.get("urls", []))

    # --- schemas (cadres 5-8) ---
    schemas = data.get("schemas", {})
    version = _version_from_path(out_path)
    if not schemas:
        warn("bloc `schemas` absent du JSON : les cadres 5-8 sont neutralises "
             "(l'exemple est retire) mais aucun contenu projet n'est injecte.")
    chrome = find_chrome() if mode == "render" else None
    if mode == "render" and chrome is None:
        mode = "placeholder"
    check_flux_consistency(schemas, flux)
    fill_schemas(doc, schemas, version, mode, chrome)
    add_schema_annex(doc, schemas, flux)
    prune_orphan_media(doc)

    doc.save(out_path)
    print(f"Ecrit: {out_path}  (schemas: mode={mode})")
    if INFO:
        print("\nInfos:")
        for m in INFO:
            print("  -", m)
    if WARN:
        print("\nAvertissements (a verifier manuellement):")
        for w in WARN:
            print("  -", w)


def _by_role(T, acteurs):
    """Ecrit directement les acteurs projet par role (lignes fixes)."""
    t = find_table(T, "Acteurs du projet")
    if t is not None:
        idx = {a["role"].upper(): a for a in acteurs if a.get("role")}
        for row in t.rows:
            cells = _unique_cells(row)
            role = ctext(cells[0]).upper()
            if role in idx and len(cells) >= 4:
                set_cell(cells[1], idx[role].get("nom"))
                set_cell(cells[2], idx[role].get("fonction"))
                set_cell(cells[3], idx[role].get("entite"))
    return []  # deja ecrit, ne pas repasser par fill_rowset


def fill_servers(T, serveurs):
    """Remplit les blocs serveurs (2 serveurs par bloc de tableau)."""
    blocks = [t for t in T if find_in_table(t, "Nom serveur (logique)")]
    slot = 0  # index de serveur logique global
    for t in blocks:
        # chaque bloc a 2 emplacements serveur (gauche/droite)
        name_row = t.rows[0]
        cells = _unique_cells(name_row)
        # positions des libelles "Nom serveur (logique)"
        pos = [i for i, c in enumerate(cells) if "Nom serveur" in ctext(c)]
        for p in pos:
            if slot >= len(serveurs):
                return
            s = serveurs[slot]
            slot += 1
            if p + 1 < len(cells):
                set_cell(cells[p + 1], s.get("nom"))
            fill_components(t, p, s.get("composants", []))


def fill_components(t, side, comps):
    """Remplit la sous-grille composants d'un cote (gauche side=0)."""
    # localise l'en-tete "Catégorie|Composant|Version|Rôle"
    hidx = None
    for i, row in enumerate(t.rows):
        if "Catégorie" in ctext(_unique_cells(row)[0]) or \
           any("Catégorie" == ctext(c) for c in _unique_cells(row)):
            hidx = i
            break
    if hidx is None:
        return
    base = 0 if side == 0 else 5  # colonnes gauche 0-3, droite ~5-8
    first = hidx + 1
    slots = len(t.rows) - first
    while slots < len(comps):
        clone_row_after(t, t.rows[first])
        slots += 1
    for j, comp in enumerate(comps):
        cells = _unique_cells(t.rows[first + j])
        vals = [comp.get("categorie"), comp.get("composant"),
                comp.get("version"), comp.get("role")]
        for k, v in enumerate(vals):
            if base + k < len(cells):
                set_cell(cells[base + k], v)


def fill_urls(T, urls):
    blocks = [t for t in T if find_in_table(t, "Libellé URL")]
    for i, t in enumerate(blocks):
        if i >= len(urls):
            break
        u = urls[i]
        rows = t.rows
        # r0: Libelle URL | <valeur>
        set_cell(_unique_cells(rows[0])[-1], u.get("libelle"))
        # r2: acteur | ressource | fonctionnalite | donnees
        if len(rows) >= 3:
            cells = _unique_cells(rows[2])
            for ci, key in enumerate(["acteur", "ressource",
                                      "fonctionnalite", "donnees"]):
                if ci < len(cells):
                    set_cell(cells[ci], u.get(key))


def find_in_table(t, anchor):
    for row in t.rows:
        for c in _unique_cells(row):
            if anchor in ctext(c):
                return True
    return False


if __name__ == "__main__":
    main()
