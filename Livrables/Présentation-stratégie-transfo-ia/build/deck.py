# -*- coding: utf-8 -*-
"""Generateur des decks 'Point d'etape IA' - charte DSFR Ministeres Sociaux.
Deux versions : epuree (oral) + detaillee (auto-portante). Icones Remix/DSFR."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")
IC = os.path.join(ASSETS, "ic")

# ---- Palette DSFR ----
INK="161616"; BODY="3A3A3A"; MENTION="6A6A6A"; BLUE="000091"; PERI="6A6AF4"
LAV="E3E3FD"; CARD="F5F5FE"; CARD2="ECECFE"; RED="E1000F"; GOLD="FFC800"
WHITE="FFFFFF"; RULE="DDDDDD"
MAR = "Marianne"
DATE = "25/06/2026"

SW, SH = 13.333, 7.5
ML = 0.85
MR = SW - 0.85

def C(h): return RGBColor.from_string(h)

def slide(prs, bg=WHITE):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    if bg:
        r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        r.fill.solid(); r.fill.fore_color.rgb = C(bg); r.line.fill.background()
        r.shadow.inherit = False
        sp = r._element; sp.getparent().remove(sp); s.shapes._spTree.insert(2, sp)
    return s

def rect(s, x, y, w, h, fill=None, line=None, lw=0.75, rounded=False, radius=0.06):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb = C(fill)
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = C(line); shp.line.width = Pt(lw)
    if rounded:
        try: shp.adjustments[0] = radius
        except Exception: pass
    return shp

def hang(p, marL=0.22):
    pPr = p._p.get_or_add_pPr()
    pPr.set('marL', str(Inches(marL))); pPr.set('indent', str(-Inches(marL)))

def tb(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    b = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = b.text_frame; tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf

def para(tf, runs, align=PP_ALIGN.LEFT, sa=4, sb=0, ls=1.06, first=False, bullet=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    if sa is not None: p.space_after = Pt(sa)
    if sb is not None: p.space_before = Pt(sb)
    if ls is not None: p.line_spacing = ls
    if bullet: hang(p)
    for rr in runs:
        r = p.add_run(); r.text = rr['t']
        f = r.font
        f.name = rr.get('font', MAR); f.size = Pt(rr.get('sz', 16))
        f.bold = rr.get('b', False); f.italic = rr.get('i', False)
        f.color.rgb = C(rr.get('c', INK))
    return p

def R(t, **k): k['t']=t; return k

def icon(s, name, x, y, size, white=False):
    suf = "_w" if white else ""
    s.shapes.add_picture(os.path.join(IC, f"{name}{suf}.png"),
                         Inches(x), Inches(y), width=Inches(size), height=Inches(size))

def node(s, name, cx, cy, dia=0.52):
    rect(s, cx-dia/2, cy-dia/2, dia, dia, fill=PERI, rounded=True, radius=0.5)
    isz = dia*0.54
    icon(s, name, cx-isz/2, cy-isz/2, isz, white=True)

def arrow_end(conn):
    ln = conn.line._get_or_add_ln()
    te = ln.makeelement(qn('a:tailEnd'), {'type':'triangle','w':'med','len':'med'})
    ln.append(te)

# ---------- Briques recurrentes ----------
def header(s):
    s.shapes.add_picture(os.path.join(ASSETS,"logo.png"), Inches(0.85), Inches(0.30), height=Inches(0.74))
    tf = tb(s, MR-4.2, 0.30, 4.2, 0.7, MSO_ANCHOR.TOP)
    para(tf, [R("Secrétariat général", sz=13.5, b=True, c=INK)], align=PP_ALIGN.RIGHT, sa=0, ls=1.05, first=True)
    para(tf, [R("Direction du numérique", sz=13.5, c=BODY)], align=PP_ALIGN.RIGHT, sa=0, ls=1.05)
    ln = s.shapes.add_connector(1, Inches(0.0), Inches(1.10), Inches(SW), Inches(1.10))
    ln.line.color.rgb = C(RULE); ln.line.width = Pt(0.75)

def footer(s, n=None):
    tf = tb(s, 0.85, SH-0.46, 6.0, 0.3)
    para(tf, [R("Ministères Sociaux – "+DATE, sz=11, c=MENTION)], sa=0, first=True)
    if n is not None:
        tf2 = tb(s, MR-0.8, SH-0.46, 0.8, 0.3)
        para(tf2, [R(str(n), sz=11, c=MENTION)], align=PP_ALIGN.RIGHT, sa=0, first=True)

def title_block(s, runs, y=1.45, w=None, x=ML):
    tf = tb(s, x, y, (w or (MR-x)), 1.4, MSO_ANCHOR.TOP)
    para(tf, runs, sa=0, ls=1.02, first=True)
    return tf

def card_bullets(tf, lst, sz=14.5):
    for it in lst:
        para(tf, [R("▪  ", sz=12, c=PERI), R(it, sz=sz, c=INK)], sa=5, ls=1.05, bullet=True)

# ============================================================
def build(detailed, out):
    D = detailed
    prs = Presentation()
    prs.slide_width = Inches(SW); prs.slide_height = Inches(SH)

    def notes(s, txt):
        s.notes_slide.notes_text_frame.text = txt

    # ---- 1. COUVERTURE ----
    s = slide(prs)
    s.shapes.add_picture(os.path.join(ASSETS,"logo.png"), Inches(0.85), Inches(0.30), height=Inches(0.74))
    tf = tb(s, MR-4.2, 0.30, 4.2, 0.7)
    para(tf, [R("Secrétariat général", sz=13.5, b=True, c=INK)], align=PP_ALIGN.RIGHT, sa=0, first=True)
    para(tf, [R("Direction du numérique", sz=13.5, c=BODY)], align=PP_ALIGN.RIGHT, sa=0)
    bx, by, bw, bh = 0.85, 1.35, 11.95, 5.75
    rect(s, bx, by, bw, bh, fill=BLUE)
    stw = 2.6; sth = stw*540/910
    s.shapes.add_picture(os.path.join(ASSETS,"stripes.png"),
                         Inches(bx+bw-stw), Inches(by+bh-sth), width=Inches(stw), height=Inches(sth))
    s.shapes.add_picture(os.path.join(ASSETS,"sparkle.png"),
                         Inches(1.55), Inches(2.70), width=Inches(1.15), height=Inches(1.15))
    tf = tb(s, 1.5, 4.20, 10.8, 1.5)
    para(tf, [R("Accompagner l’adoption", sz=44, b=True, c=WHITE)], sa=0, ls=1.0, first=True)
    para(tf, [R("de l’IA", sz=44, b=True, c=WHITE)], sa=0, ls=1.0)
    tf = tb(s, 1.53, 5.85, 10.4, 0.7)
    para(tf, [R("Point d’étape — enjeux, exploration et plan d’action", sz=20, c=WHITE)], sa=0, first=True)
    notes(s, "Point d'etape de la mission IA a la Direction du numerique des Ministeres Sociaux. "
             "Objectif du jour : partager ou en est la mission - enjeux, exploration menee, plan d'action sur 3 horizons. "
             "10 min de presentation puis echange.")

    # ---- 2. SOMMAIRE ----
    s = slide(prs); header(s); footer(s, 2)
    title_block(s, [R("Au programme", sz=33, b=True, c=INK)])
    items = [
        ("Les enjeux de la mission", "Ce que je dois apporter aux équipes du numérique."),
        ("La phase d’exploration", "Comprendre le terrain et observer les usages de l’IA."),
        ("Le plan d’action", "Trois horizons : court, moyen et long terme."),
    ]
    y = 2.7
    for i,(t,d) in enumerate(items, 1):
        rect(s, ML, y, 0.62, 0.62, fill=PERI, rounded=True, radius=0.5)
        tfn = tb(s, ML, y, 0.62, 0.62, MSO_ANCHOR.MIDDLE)
        para(tfn, [R(str(i), sz=22, b=True, c=WHITE)], align=PP_ALIGN.CENTER, sa=0, first=True)
        tft = tb(s, ML+0.95, y-0.06, 10.0, 0.8, MSO_ANCHOR.MIDDLE if not D else MSO_ANCHOR.TOP)
        para(tft, [R(t, sz=21, b=True, c=INK)], sa=0, first=True)
        if D: para(tft, [R(d, sz=14.5, c=BODY)], sa=0, sb=3)
        y += 1.18 if not D else 1.28
    notes(s, "Trois temps : les enjeux (mon mandat), l'exploration (ce que j'ai fait pour comprendre), "
             "puis le plan d'action structure en trois horizons.")

    # ---- DIVIDER ----
    def divider(num, plain, rest, note):
        s = slide(prs); header(s); footer(s, None)
        tf = tb(s, ML, 1.3, MR-ML, 5.4, MSO_ANCHOR.MIDDLE)
        para(tf, [R(num, sz=20, b=True, c=PERI)], align=PP_ALIGN.CENTER, sa=14, first=True)
        para(tf, [R(plain+rest, sz=46, b=True, c=INK)], align=PP_ALIGN.CENTER, sa=0, ls=1.0)
        notes(s, note)

    # ---- 3. DIVIDER Enjeux ----
    divider("01", "Les ", "enjeux", "Mon mandat sur cette mission, en trois enjeux.")

    # ---- 4. ENJEUX ----
    s = slide(prs); header(s); footer(s, 4)
    title_block(s, [R("Les enjeux", sz=33, b=True, c=INK)])
    tf = tb(s, ML, 2.18, MR-ML, 0.5)
    para(tf, [R("Faire de l’IA un levier concret pour les métiers du numérique.", sz=16.5, c=BODY)], sa=0, first=True)
    cards = [
        ("search-eye","Identifier les usages","Repérer les cas d’usage IA pour chaque métier.",
         "Architectes, chefs de projet, designers, développeurs."),
        ("user-follow","Accompagner la maîtrise","Aider les équipes à s’approprier des cas concrets.",
         "Passer de l’essai isolé à une pratique maîtrisée."),
        ("megaphone","Évangéliser","Diffuser la culture et les bons réflexes de l’IA.",
         "Créer l’adhésion et l’envie d’essayer."),
    ]
    cw = (MR-ML-2*0.45)/3; cy=2.9; chh=2.75 if not D else 2.95
    for i,(ic,t,d,extra) in enumerate(cards):
        cx = ML + i*(cw+0.45)
        rect(s, cx, cy, cw, chh, fill=CARD2, rounded=True, radius=0.05)
        icon(s, ic, cx+0.34, cy+0.34, 0.6)
        tft = tb(s, cx+0.34, cy+1.14, cw-0.64, chh-1.3)
        para(tft, [R(t, sz=18.5, b=True, c=BLUE)], sa=6, ls=1.02, first=True)
        para(tft, [R(d, sz=14.5, c=BODY)], sa=0, ls=1.1)
        if D: para(tft, [R(extra, sz=13, c=MENTION, i=True)], sa=0, sb=6, ls=1.08)
    fy = cy+chh+0.32
    tf = tb(s, ML, fy, 2.2, 0.5, MSO_ANCHOR.MIDDLE)
    para(tf, [R("Pour 4 métiers :", sz=14, b=True, c=INK)], sa=0, first=True)
    rx = ML+2.25
    for rname in ["Architectes","Chefs de projet","Designers","Développeurs"]:
        wch = 0.30 + len(rname)*0.105
        rect(s, rx, fy+0.02, wch, 0.46, fill=None, line=PERI, lw=1.0, rounded=True, radius=0.5)
        tfc = tb(s, rx, fy+0.02, wch, 0.46, MSO_ANCHOR.MIDDLE)
        para(tfc, [R(rname, sz=13, b=True, c=BLUE)], align=PP_ALIGN.CENTER, sa=0, first=True)
        rx += wch+0.22
    notes(s, "Mon mandat tient en trois enjeux : identifier les usages pertinents par metier, accompagner "
             "la montee en maitrise sur des cas concrets, et evangeliser pour diffuser la culture IA. "
             "Au service de quatre metiers : architectes, chefs de projet, designers, developpeurs.")

    # ---- 5. DIVIDER Exploration ----
    divider("02", "L’", "exploration", "Comment j'ai pris la mesure du terrain et des usages en place.")

    # ---- 6. EXPLORATION : rencontres ----
    s = slide(prs); header(s); footer(s, 6)
    title_block(s, [R("Aller à la rencontre du terrain", sz=33, b=True, c=INK)])
    blocks = [
        ("team","Les équipes", ["SRDT","DACCORD","Egapro","DomiFA","SIRENA"],
         "Comprendre leur quotidien et leurs enjeux.",
         "Garder un accès à leur contexte à la demande."),
        ("briefcase","Les métiers", ["Architectes","Products","Designers","Développeurs"],
         "Cerner les besoins propres à chaque rôle.", ""),
        ("discuss","Échanges clés", ["Dominique — RGAA / accessibilité","Victor Degliame — outillage interne"],
         "Sécuriser conformité et stratégie d’outillage.", ""),
    ]
    cw = (MR-ML-2*0.45)/3; cy=2.35; chh=3.95 if not D else 4.4
    for i,(ic,t,lst,d,extra) in enumerate(blocks):
        cx = ML + i*(cw+0.45)
        rect(s, cx, cy, cw, chh, fill=CARD2, rounded=True, radius=0.05)
        icon(s, ic, cx+0.34, cy+0.34, 0.56)
        tft = tb(s, cx+0.34, cy+1.05, cw-0.68, chh-1.2)
        para(tft, [R(t, sz=18, b=True, c=BLUE)], sa=9, ls=1.0, first=True)
        card_bullets(tft, lst)
        para(tft, [R(d, sz=13.5, c=BODY, i=True)], sa=0, sb=9, ls=1.1)
        if D and extra: para(tft, [R(extra, sz=12.5, c=MENTION)], sa=0, sb=6, ls=1.08)
    notes(s, "Premier reflexe : aller voir les gens. J'ai rencontre cinq equipes produit (SRDT, DACCORD, "
             "Egapro, DomiFA, SIRENA) pour comprendre leur quotidien et garder un acces au contexte. "
             "J'ai croise les quatre metiers, et securise deux sujets cles : l'accessibilite/RGAA avec Dominique, "
             "et la strategie d'outillage des internes avec Victor Degliame.")

    # ---- 7. EXPLORATION : Egapro (observer / proposer) ----
    s = slide(prs); header(s); footer(s, 7)
    title_block(s, [R("Egapro : un usage de l’IA déjà avancé", sz=31, b=True, c=INK)])
    tf = tb(s, ML, 2.2, MR-ML, 0.6)
    para(tf, [R("Une équipe qui exploite déjà Claude — j’explore cet usage en développant dessus.",
                sz=16.5, c=BODY)], sa=0, ls=1.1, first=True)
    cw = (MR-ML-0.5)/2; cy=3.05; chh=3.25
    # gauche : observer
    rect(s, ML, cy, cw, chh, fill=CARD2, rounded=True, radius=0.05)
    icon(s, "terminal-box", ML+0.36, cy+0.34, 0.56)
    tl = tb(s, ML+0.36, cy+1.05, cw-0.72, chh-1.2)
    para(tl, [R("Ce que j’observe", sz=18, b=True, c=BLUE)], sa=10, first=True)
    obs = ["Un usage de Claude déjà en place avant mon arrivée.", "Une orchestration ambitieuse."]
    if D: obs.append("Une base concrète pour mesurer le potentiel réel.")
    card_bullets(tl, obs, sz=15)
    # droite : proposer
    rx = ML+cw+0.5
    rect(s, rx, cy, cw, chh, fill=CARD2, rounded=True, radius=0.05)
    icon(s, "settings-3", rx+0.36, cy+0.34, 0.56)
    tr = tb(s, rx+0.36, cy+1.05, cw-0.72, chh-1.2)
    para(tr, [R("Ce que je propose", sz=18, b=True, c=BLUE)], sa=10, first=True)
    prop = ["Des optimisations sur la façon dont l’IA est mobilisée.",
            "Intégrer des estimations par taille de « T-shirt » — aujourd’hui en expérimentation."]
    card_bullets(tr, prop, sz=15)
    notes(s, "Egapro est une des equipes : Claude y est DEJA fortement exploite, avec une orchestration "
             "ambitieuse mise en place avant mon arrivee. Mon role est d'explorer cet usage en developpant "
             "dessus. Ce que j'observe : un usage avance deja en place. Ce que je propose a partir de la : "
             "des optimisations, et l'integration d'estimations par taille de T-shirt (aujourd'hui en experimentation). "
             "C'est de l'exploration et des pistes, pas un livrable abouti.")

    # ---- 8. DIVIDER Plan d'action ----
    divider("03", "Le ", "plan d’action", "Ce que je propose, sequence du court au long terme.")

    # ---- 9. ROADMAP : timeline 3 horizons ----
    s = slide(prs); header(s); footer(s, 9)
    title_block(s, [R("Trois horizons", sz=33, b=True, c=INK)])
    tf = tb(s, ML, 2.16, MR-ML, 0.45)
    para(tf, [R("Une trajectoire progressive, du quick win à la cible souveraine.", sz=16, c=BODY)], sa=0, first=True)
    hor = [
        ("rocket","COURT TERME","Outiller chaque métier",
         ["Architectes","Chefs de projet","Designers","Développeurs"]),
        ("settings-3","MOYEN TERME","Industrialiser & animer",
         ["Communauté de champions","CI/CD & déploiement","Dev agentique","Observabilité"]),
        ("government","LONG TERME","Un harness maison",
         ["Souveraineté","Conformité intégrée","Expérience maîtrisée"]),
    ]
    n=3; gap=0.62
    cw=(MR-ML-(n-1)*gap)/n
    cxs=[ML+i*(cw+gap) for i in range(n)]
    centers=[cx+cw/2 for cx in cxs]
    axis_y=3.32
    ln = s.shapes.add_connector(1, Inches(centers[0]), Inches(axis_y), Inches(centers[-1]+0.55), Inches(axis_y))
    ln.line.color.rgb=C(PERI); ln.line.width=Pt(1.75); arrow_end(ln)
    for i,(ic,lab,t,lst) in enumerate(hor):
        tfl = tb(s, cxs[i], 2.78, cw, 0.34, MSO_ANCHOR.MIDDLE)
        para(tfl, [R(lab, sz=13, b=True, c=PERI)], align=PP_ALIGN.CENTER, sa=0, first=True)
        node(s, ic, centers[i], axis_y, dia=0.52)
    cy=3.78; chh=2.5
    for i,(ic,lab,t,lst) in enumerate(hor):
        rect(s, cxs[i], cy, cw, chh, fill=CARD2, rounded=True, radius=0.05)
        tft = tb(s, cxs[i]+0.32, cy+0.3, cw-0.64, chh-0.5)
        para(tft, [R(t, sz=18, b=True, c=BLUE)], sa=10, ls=1.0, first=True)
        card_bullets(tft, lst, sz=14)
    notes(s, "Le plan se lit sur trois horizons. Court terme : des quick wins par metier. Moyen terme : "
             "industrialiser (CI/CD, deploiement, dev agentique, observabilite) et animer une communaute de "
             "champions. Long terme : un harness maison pour la souverainete et la conformite. Je detaille ensuite.")

    # ---- 10. COURT TERME ----
    s = slide(prs); header(s); footer(s, 10)
    title_block(s, [R("Court terme : outiller chaque métier", sz=31, b=True, c=INK)])
    ct = [
        ("compasses-2","Architectes","Générer des Dossiers d’Architecture (DA) avec l’IA.",""),
        ("clipboard","Chefs de projet","Pré-audit d’homologation continu et meilleur pilotage des projets faits avec l’IA.",
         "Rencontrer plus de CDP ; expérimenter les estimations « T-shirt » sur Egapro."),
        ("palette","Designers","Créer et utiliser des skills d’UX et d’audit UX pour générer et challenger des prototypes rapidement.",
         "Puis analyser les autres besoins existants."),
        ("code-s-slash","Développeurs","Aller vers ceux en difficulté avec l’IA et les guider.",""),
    ]
    cw = (MR-ML-0.45)/2; chh=1.7 if not D else 2.02; cy=2.5; gap=0.3
    for i,(ic,t,d,extra) in enumerate(ct):
        col=i%2; row=i//2
        cx=ML+col*(cw+0.45); ry=cy+row*(chh+gap)
        rect(s, cx, ry, cw, chh, fill=CARD2, rounded=True, radius=0.06)
        icon(s, ic, cx+0.34, ry+0.30, 0.5)
        tfh = tb(s, cx+1.0, ry+0.30, cw-1.3, 0.5, MSO_ANCHOR.MIDDLE)
        para(tfh, [R(t, sz=17.5, b=True, c=BLUE)], sa=0, first=True)
        tft = tb(s, cx+0.34, ry+0.92, cw-0.68, chh-1.05)
        para(tft, [R(d, sz=14, c=BODY)], sa=0, ls=1.08, first=True)
        if D and extra: para(tft, [R(extra, sz=12.5, c=MENTION, i=True)], sa=0, sb=4, ls=1.05)
    notes(s, "Court terme, metier par metier. Architectes : generer les dossiers d'architecture avec l'IA. "
             "Chefs de projet : un pre-audit d'homologation continu et l'enjeu de pilotage des projets faits "
             "avec l'IA (estimations T-shirt experimentees sur Egapro). Designers : des skills d'UX et d'audit "
             "pour prototyper et challenger vite. Developpeurs : accompagner ceux qui peinent a s'approprier l'IA.")

    # ---- 11. MOYEN TERME ----
    s = slide(prs); header(s); footer(s, 11)
    title_block(s, [R("Moyen terme : industrialiser & animer", sz=31, b=True, c=INK)])
    lw_ = 3.85; cy=2.55; chh=3.7
    rect(s, ML, cy, lw_, chh, fill=BLUE, rounded=True, radius=0.05)
    tft = tb(s, ML+0.4, cy+0.45, lw_-0.8, chh-0.9, MSO_ANCHOR.TOP)
    para(tft, [R("Animer la communauté", sz=20, b=True, c=WHITE)], sa=8, ls=1.02, first=True)
    para(tft, [R("Identifier et préparer l’animation d’une communauté de champions pour chaque métier.",
                 sz=15, c="E3E3FD")], sa=0, ls=1.18)
    if D:
        para(tft, [R("Des relais qui diffusent les pratiques au plus près des équipes.", sz=13, c="CFCFFB", i=True)], sa=0, sb=10, ls=1.1)
    rx = ML+lw_+0.5; rw = MR-rx
    tf = tb(s, rx, cy+0.02, rw, 0.5)
    para(tf, [R("Volet développeurs", sz=18, b=True, c=BLUE)], sa=0, first=True)
    devs = [
        ("CI/CD", "bots de revue de code automatiques."),
        ("Déploiement", "dès les licences Claude internes (Sièges)."),
        ("Dev agentique", "pour les internes, en lien avec la DINUM."),
        ("Observabilité", "suivre les usages et la qualité."),
    ]
    yy = cy+0.66
    for t,d in devs:
        rect(s, rx, yy+0.06, 0.15, 0.15, fill=PERI, rounded=True, radius=0.5)
        tft = tb(s, rx+0.32, yy, rw-0.32, 0.6)
        para(tft, [R(t+" — ", sz=15, b=True, c=INK), R(d, sz=14.5, c=BODY)], sa=0, ls=1.05, first=True)
        yy += 0.84
    notes(s, "Moyen terme, deux volets. Animer : identifier et faire vivre une communaute de champions par "
             "metier. Cote developpeurs : fiabiliser la CI/CD avec des bots de review, activer le deploiement "
             "des qu'on aura les licences Claude internes, ouvrir le dev agentique avec la DINUM, mettre en "
             "place l'observabilite.")

    # ---- 12. LONG TERME : harness ----
    s = slide(prs); header(s); footer(s, 12)
    title_block(s, [R("Long terme : un harness maison", sz=31, b=True, c=INK)])
    tf = tb(s, ML, 2.18, MR-ML, 0.5)
    para(tf, [R("Construire notre propre environnement IA, maîtrisé de bout en bout.", sz=16.5, c=BODY)], sa=0, first=True)
    ben = [
        ("user-smile","Expérience maîtrisée","Embarquer les profils les moins technophiles."),
        ("focus-3","Périmètre maîtrisé","Contrôler les besoins auxquels l’outil répond."),
        ("shield-check","Conformité intégrée","Intégrer accessibilité et conformité dans les réalisations."),
        ("bank","Souveraineté & stabilité","Indépendance vis-à-vis des fournisseurs."),
    ]
    cw=(MR-ML-3*0.4)/4; cy=2.95; chh=2.65
    for i,(ic,t,d) in enumerate(ben):
        cx=ML+i*(cw+0.4)
        rect(s, cx, cy, cw, chh, fill=CARD2, rounded=True, radius=0.06)
        icon(s, ic, cx+0.3, cy+0.34, 0.56)
        tft=tb(s, cx+0.3, cy+1.1, cw-0.56, chh-1.25)
        para(tft, [R(t, sz=16, b=True, c=BLUE)], sa=8, ls=1.02, first=True)
        para(tft, [R(d, sz=13.5, c=BODY)], sa=0, ls=1.12)
    notes(s, "L'horizon cible : un harness maison, notre propre environnement IA. Quatre benefices : une "
             "experience maitrisee pour embarquer les moins technophiles, un perimetre maitrise sur les besoins "
             "couverts, la conformite (accessibilite) integree des la conception, et la souverainete - "
             "l'independance vis-a-vis des fournisseurs.")

    # ---- 13. SYNTHESE ----
    s = slide(prs); header(s); footer(s, 13)
    title_block(s, [R("Où nous en sommes", sz=33, b=True, c=INK)])
    colw=(MR-ML-0.5)/2; cy=2.4
    tf=tb(s, ML, cy, colw, 0.5)
    para(tf, [R("Ce qui est acquis", sz=18, b=True, c=BLUE)], sa=0, first=True)
    acq=[ "Cinq équipes rencontrées : SRDT, DACCORD, Egapro, DomiFA, SIRENA.",
          "Métiers rencontrés : architectes, products, designers, développeurs.",
          "Échanges menés sur l’accessibilité/RGAA (Dominique) et l’outillage interne (V. Degliame).",
          "Egapro : usage avancé de Claude observé, premières pistes proposées.",
          "Un plan d’action structuré sur trois horizons." ]
    tfa=tb(s, ML, cy+0.58, colw, 4.2)
    for it in acq:
        para(tfa, [R("▪  ", sz=12, c=PERI), R(it, sz=15, c=BODY)], sa=8, ls=1.1, bullet=True, first=(it==acq[0]))
    rx=ML+colw+0.5
    tf=tb(s, rx, cy, colw, 0.5)
    para(tf, [R("Prochaines étapes", sz=18, b=True, c=BLUE)], sa=0, first=True)
    nxt=[ "Architectes : outiller la génération de DA avec l’IA.",
          "Chefs de projet : en rencontrer plus, lancer le pré-audit d’homologation continu.",
          "Designers : créer les premiers skills d’UX et d’audit UX, analyser les autres besoins.",
          "Développeurs : aller vers ceux en difficulté avec l’IA et les guider.",
          "Préparer la communauté de champions ; anticiper licences Claude & lien DINUM." ]
    tfn=tb(s, rx, cy+0.58, colw, 4.2)
    for it in nxt:
        para(tfn, [R("▪  ", sz=12, c=PERI), R(it, sz=15, c=BODY)], sa=8, ls=1.1, bullet=True, first=(it==nxt[0]))
    notes(s, "Synthese. Ce qui est acquis : les cinq equipes rencontrees (SRDT, DACCORD, Egapro, DomiFA, SIRENA), "
             "les metiers rencontres (architectes, products, designers, developpeurs), les echanges sur le RGAA "
             "avec Dominique et l'outillage interne avec Victor Degliame, l'usage avance de Claude observe sur "
             "Egapro avec de premieres pistes, et un plan structure sur 3 horizons. Prochaines etapes, surtout "
             "court terme : DA pour les architectes ; rencontrer plus de CDP et lancer le pre-audit d'homologation "
             "continu ; premiers skills UX/audit cote designers ; accompagner les developpeurs en difficulte. "
             "Puis preparer la communaute de champions et anticiper les licences Claude et le lien DINUM. "
             "Place a vos questions.")

    # ---- 14. CLOTURE ----
    s = slide(prs)
    s.shapes.add_picture(os.path.join(ASSETS,"logo.png"), Inches(0.85), Inches(0.30), height=Inches(0.74))
    rect(s, bx, by, bw, bh, fill=BLUE)
    stw=3.1; sth=stw*540/910
    s.shapes.add_picture(os.path.join(ASSETS,"stripes.png"),
                         Inches(bx+bw-stw), Inches(by+bh-sth), width=Inches(stw), height=Inches(sth))
    tf=tb(s, 1.5, 3.4, 9.6, 1.3, MSO_ANCHOR.TOP)
    para(tf, [R("Place à l’échange", sz=50, b=True, c=WHITE)], sa=0, ls=1.0, first=True)
    tf=tb(s, 1.53, 4.5, 9.8, 0.7)
    para(tf, [R("Vos questions, vos retours, vos arbitrages.", sz=22, c=WHITE)], sa=0, first=True)
    notes(s, "Transition vers les 35 min d'echange : questions, retours, arbitrages eventuels "
             "(licences, lien DINUM, priorisation des horizons).")

    prs.save(out)
    print("OK ->", os.path.basename(out), "|", len(list(prs.slides)), "slides")

if __name__ == "__main__":
    build(False, os.path.join(HERE, "IA_Point-etape_oral.pptx"))
    build(True,  os.path.join(HERE, "IA_Point-etape_detaille.pptx"))
